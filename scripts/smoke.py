"""Exercise a running local app + actual Mermaid renderer without a browser UI."""
import io
import sys
import time
import zipfile
import os
from pathlib import Path
import httpx
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image

base = sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8000'
client = httpx.Client(base_url=base, headers={'X-Workspace-Request':'1'}, timeout=60, trust_env=False)


def request(method, path, **kwargs):
    response = client.request(method,path,**kwargs)
    response.raise_for_status()
    return response


def wait_for_job(pid):
    start=time.monotonic()
    previous=''
    while time.monotonic()-start<180:
        detail=request('GET','/api/projects/'+pid).json()
        job=detail['jobs'][-1]
        if job['status']!=previous:
            print('Job:',job['status'],flush=True)
            previous=job['status']
        if job['status']=='Failed':raise AssertionError(job['error'])
        if job['status']=='Completed':return detail
        time.sleep(1)
    raise AssertionError('Timed out waiting for the worker')


pid=None
try:
    assert request('GET','/api/config').json()['demo'], 'Smoke script is for local demo instances only'
    assert request('GET','/').status_code==200
    request('POST','/api/auth/demo')
    pid=request('POST','/api/projects',json={'title':'Automated smoke validation'}).json()['id']
    request('POST',f'/api/projects/{pid}/attachments',files={'file':('requirements.md',b'# Identity\nEnterprise SSO required.\n# Recovery\nRPO of 24 hours is a supplied requirement.')})
    request('POST',f'/api/projects/{pid}/generate',json={'prompt':'Design an enterprise knowledge portal with SSO and recovery. Include architecture and estimates.','preferences':{'cloud':'AWS','outputs':['solution','diagrams','estimates','plan','presentation']}})
    detail=wait_for_job(pid)
    v=detail['versions'][-1]
    failed=[a for a in v['artifacts'] if a['status']!='Completed']
    assert not failed, failed
    assert v['context_sources'] and v['totals']['effort']['expected']==52
    for artifact in v['artifacts']:
        url=request('POST',f'/api/artifacts/{artifact["id"]}/download-link').json()['url']
        data=request('GET',url).content
        ext=artifact['filename'].split('.')[-1]
        if ext=='docx':assert Document(io.BytesIO(data)).paragraphs
        elif ext=='pdf':assert len(PdfReader(io.BytesIO(data)).pages)>0
        elif ext=='xlsx':assert load_workbook(io.BytesIO(data)).active['D9'].value=='=SUM(D4:D8)'
        elif ext=='pptx':
            presentation=Presentation(io.BytesIO(data))
            headings='\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
            for heading in ['Executive summary','Proposed solution','Estimation overview','Delivery roadmap','Assumptions for review']:
                assert heading in headings
            assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)
            assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes)
            if os.environ.get('SMOKE_CAPTURE_DIR'):
                capture=Path(os.environ['SMOKE_CAPTURE_DIR'])
                capture.mkdir(parents=True,exist_ok=True)
                (capture/artifact['filename']).write_bytes(data)
        elif ext in ['png','jpg']:
            img=Image.open(io.BytesIO(data));assert img.width>100 and img.height>100
        elif ext=='svg':assert b'<svg' in data
        elif ext=='mmd':assert b'flowchart LR' in data
        elif ext=='zip':
            with zipfile.ZipFile(io.BytesIO(data)) as z:assert len(z.namelist())==9
        print('Verified:',artifact['filename'],len(data),'bytes',flush=True)
    request('POST',f'/api/projects/{pid}/generate',json={'prompt':'Refine the previous design for stronger disaster recovery.'})
    updated=wait_for_job(pid)
    assert len(updated['versions'])==2 and updated['versions'][0]['id']==v['id']
    print('PASS: reference extraction, durable generation, actual diagrams, all document formats, download package and follow-up versions.',flush=True)
finally:
    if pid:
        try:
            response=client.delete('/api/projects/'+pid)
            print('Smoke project cleanup:',response.status_code,flush=True)
        except httpx.HTTPError:
            print('Smoke cleanup could not reach the API; test project:',pid,flush=True)
    client.close()
