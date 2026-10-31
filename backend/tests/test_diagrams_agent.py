import io
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.diagrams import DiagramsAgent, ArchitectureViews, CORE, DELIVERY, delivery_views
from app import diagrams
from app.config import settings
from app.artifacts import create_document
from test_workspace import sample_solution, fake_render
from test_presentation import visible_text


def context():
    return {'preferences':{'outputs':['diagrams','presentation']}}


def test_agent_core_types_delivery_values_and_proposed_racii():
    s = sample_solution()
    DiagramsAgent().run(context(), s)
    assert CORE | DELIVERY == {d.kind for d in s.diagrams}
    assert len(s.diagrams) == 9
    assert all(len(d.description) >= 160 for d in s.diagrams)
    assert all(row.accountable == 'Delivery lead (proposed)' for row in s.racii)
    assert [r.activity for r in s.racii] == [e.activity for e in s.estimates]
    view = next(d for d in s.diagrams if d.kind == 'estimations')
    from app.schemas import calculate
    effort=calculate(s)['effort']
    assert f'{effort["low"]:g} / {effort["expected"]:g} / {effort["high"]:g}' in view.source
    s.estimates[0].expected = s.estimates[0].high
    updated, _ = delivery_views(s)
    effort=calculate(s)['effort']
    assert f'{effort["low"]:g} / {effort["expected"]:g} / {effort["high"]:g}' in next(d for d in updated if d.kind == 'estimations').source


def test_live_agent_rejects_missing_core_views_without_demo_fallback(monkeypatch):
    s = sample_solution()
    monkeypatch.setattr(settings, 'llm_provider', 'bedrock')
    def structured(model, system, payload, validate):
        assert model is ArchitectureViews and 'network' in system and '100-160 words' in system
        good = DiagramsAgent().demo(s)
        good.diagrams[0].kind = 'integration'
        return validate(good)
    monkeypatch.setattr(diagrams, 'structured', structured)
    with pytest.raises(ValueError, match='network, conceptual, technical and logical'):
        DiagramsAgent().run(context(), s)


def test_each_diagram_has_own_slide_description_delivery_section_and_matrix():
    s = sample_solution()
    DiagramsAgent().run(context(), s)
    s.diagrams[0].description += ' Detailed considerations remain essential for implementation.' * 25
    images = [(d.title, fake_render(d.source)['png']) for d in s.diagrams]
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'demo', images)))
    pictures = []
    for slide in prs.slides:
        images = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(images) <= 1
        if images:
            pictures.append(slide)
            assert 'How to read this diagram' in '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        for shape in slide.shapes:
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
    assert len(pictures) == len(s.diagrams)
    text = visible_text(prs)
    assert 'Delivery' in text and 'RACII Matrix' in text and 'I2: Involved' in text
    assert 'detailed description' in text
    assert text.count('Detailed considerations remain essential for implementation.') == 25
    tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert any(t.cell(0, 5).text == 'I2' for t in tables if len(t.columns) == 6)
