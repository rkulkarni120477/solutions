import io
import zipfile
from decimal import Decimal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app import worker
from app.artifacts import create_document
from app.providers import demo
from app.schemas import Preferences, calculate
from app import artifacts, providers
from app.db import Artifact, Version
from app.config import settings
from test_workspace import fake_render


def solution():
    return demo({'preferences': Preferences().model_dump(), 'current_prompt': 'Design a knowledge portal with SSO and AI search',
                 'approved_decisions': [], 'history': [], 'reference_passages': []})


def visible_text(prs):
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return '\n'.join(parts)


def test_briefing_coverage_native_estimates_costs_and_bounds():
    s = solution()
    for row in s.estimates:
        row.rate = 800
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'demo', metadata={'version': 7})))
    text = visible_text(prs)
    for phrase in ['Executive summary', 'Proposed solution', 'Estimation overview', 'Effort by workstream',
                   'Implementation cost', 'Operating costs and exclusions', 'Delivery roadmap',
                   'Assumptions for review', 'Open questions and decisions', 'Response version 7', 'DEMO:']:
        assert phrase in text
    totals = calculate(s)
    assert f"USD {totals['implementation_cost']['expected']:,.0f}" in text
    expected = totals['with_contingency']['expected']
    assert f'{expected:,.2f}'.rstrip('0').rstrip('.') in text
    tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert any(table.cell(0, 0).text == 'Activity' and table.cell(0, 3).text == 'Expected' for table in tables)
    charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    assert list(charts[0].series[0].values) == [row.expected for row in s.estimates]
    assert prs.slide_width / prs.slide_height > 1.77
    assert 'Diagram export is unavailable' in text
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10


def test_unpriced_and_unrequested_estimates_are_explicit():
    s = solution()
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'openai')))
    assert 'Supply a daily rate for every workstream' in visible_text(prs)
    s.estimates, s.phases, s.diagrams = [], [], []
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'demo')))
    text = visible_text(prs)
    assert 'No effort estimate was requested or provided' in text
    assert 'No implementation plan was requested or provided' in text
    assert 'No diagram was requested or provided' in text
    assert not any(shape.has_chart for slide in prs.slides for shape in slide.shapes)


def test_presentation_pagination_preserves_long_text_and_table_rows():
    s = solution()
    sentence = 'Reference detail with required security and delivery considerations.'
    s.sections[1].content = ' '.join([sentence] * 80)
    for row in s.estimates:
        row.activity = 'Discovery and architecture planning with all required stakeholder dependencies ' * 2
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'demo')))
    text = visible_text(prs)
    assert text.count(sentence) == 80
    assert 'continued 2' in text
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= prs.slide_height + 10


def test_package_automatically_includes_pptx_and_refresh_preserves_version(workspace, monkeypatch):
    alice, bob, _ = workspace
    monkeypatch.setattr(worker, 'render_diagram', fake_render)
    monkeypatch.setattr(artifacts, 'render_diagram', fake_render)
    pid = alice.post('/api/projects', json={'title': 'Briefing package'}).json()['id']
    # Explicitly omit presentation, diagrams and solution-document flags.
    alice.post(f'/api/projects/{pid}/generate', json={'prompt': 'Estimate a secure customer portal', 'preferences': {'outputs': ['estimates']}})
    worker.run_once()
    original = alice.get(f'/api/projects/{pid}').json()['versions'][0]
    assert {a['kind'] for a in original['artifacts']} == {'pptx', 'xlsx', 'zip',*[f'diagram:0:{ext}' for ext in ['mmd','svg','png','jpg']]}
    package = next(a for a in original['artifacts'] if a['kind'] == 'zip')
    link = alice.post(f'/api/artifacts/{package["id"]}/download-link').json()['url']
    with zipfile.ZipFile(io.BytesIO(alice.get(link).content)) as archive:
        pptx_name = next(name for name in archive.namelist() if name.endswith('.pptx'))
        assert 'Executive summary' in visible_text(Presentation(io.BytesIO(archive.read(pptx_name))))
    assert bob.post(f'/api/projects/{pid}/versions/{original["id"]}/refresh-files').status_code == 404
    def no_provider(*args, **kwargs):
        raise AssertionError('Refreshing files must not invoke an LLM')
    monkeypatch.setattr(worker, 'generate', no_provider)
    response = alice.post(f'/api/projects/{pid}/versions/{original["id"]}/refresh-files')
    assert response.status_code == 202
    worker.run_once()
    versions = alice.get(f'/api/projects/{pid}').json()['versions']
    assert len(versions) == 2
    assert versions[0]['artifacts'] == original['artifacts']
    assert versions[1]['response'] == original['response']
    assert versions[1]['provider'] == original['provider']
    assert all(a['status'] == 'Completed' for a in versions[1]['artifacts'])


def test_presentation_only_includes_architecture_and_estimate_inputs():
    s = demo({'preferences': Preferences(outputs=['presentation']).model_dump(),
              'current_prompt': 'Design a secure knowledge portal', 'reference_passages': [],
              'approved_decisions': [], 'history': []})
    assert any(d.kind == 'logical' for d in s.diagrams) and s.estimates
    diagrams = [(d.title, fake_render(d.source)['png']) for d in s.diagrams]
    prs = Presentation(io.BytesIO(create_document('pptx', s, 'demo', diagrams, {'version': 3})))
    for slide in prs.slides:
        assert any(shape.name == 'Project header' and shape.text == s.title for shape in slide.shapes)
        assert any(shape.name == 'Review footer' and 'v3 |' in shape.text for shape in slide.shapes)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in prs.slides for shape in slide.shapes)
    assert 'Effort by workstream' in visible_text(prs)
    assert next(i for i,slide in enumerate(prs.slides) if any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)) < 4


def test_ppt_retry_renders_architecture_without_sibling_png(workspace, monkeypatch):
    alice, _, sessions = workspace
    pid = alice.post('/api/projects', json={'title': 'PPT retry diagram'}).json()['id']
    s = solution()
    calls = []
    def render(source):
        calls.append(source)
        return fake_render(source)
    monkeypatch.setattr(artifacts, 'render_diagram', render)
    with sessions() as db:
        version = Version(project_id=pid, number=1, prompt='PPT retry', preferences={}, response=s.model_dump(), provider='demo')
        db.add(version)
        db.flush()
        artifact = Artifact(project_id=pid, version_id=version.id, kind='pptx', filename='retry.pptx')
        db.add(artifact)
        db.flush()
        data = artifacts.build_artifact(db, artifact, version, {})
    prs = Presentation(io.BytesIO(data))
    assert calls == [d.source for d in s.diagrams]
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in prs.slides for shape in slide.shapes)


def test_initial_solution_defers_technical_artifacts_to_specialist_agents(monkeypatch):
    monkeypatch.setattr(settings, 'llm_provider', 'bedrock')
    monkeypatch.setattr(settings, 'llm_model', 'test')
    monkeypatch.setattr(providers.time, 'sleep', lambda _: None)
    context = {'preferences': Preferences(outputs=['presentation']).model_dump(), 'current_prompt': 'Design a portal',
               'reference_passages': [], 'history': [], 'approved_decisions': []}
    s = solution()
    missing = s.model_copy(deep=True)
    missing.diagrams, missing.estimates = [], []
    results = iter([missing.model_dump_json(), s.model_dump_json()])
    monkeypatch.setattr(providers, '_complete', lambda *args: next(results))
    result = providers.generate(context)
    assert result.diagrams == [] and result.estimates == []
