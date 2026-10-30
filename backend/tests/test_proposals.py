import io
import pytest
from docx import Document
from pptx import Presentation
from app import worker, artifacts
from app.db import Proposal, AgentMessage, ResponseDocument, Artifact
from test_workspace import fake_render

VAGUE = 'We need a proposal for a new customer portal for our insurance business.'
COMPLETE = ('Client: Contoso Insurance. Objective: reduce claim handling time by 30%. Scope: build a customer self-service portal with claims '
            'submission and integration to the policy system. Timeline: go-live within 4 months. Budget: fixed price under USD 400k. '
            'Platform: Azure with .NET and React. Scale: 20,000 users, 99.9% availability and strong security. The portal must support SSO.')


def test_live_agent_failure_never_uses_demo_output(monkeypatch):
    from app import agents
    from app.providers import ProviderError
    from app.schemas import ProposalPlan
    messages = []
    class Bus:
        def send(self, *message):
            messages.append(message)
    monkeypatch.setattr(agents.settings, 'llm_provider', 'bedrock')
    def unavailable(*args):
        raise ProviderError('Model access denied')
    monkeypatch.setattr(agents, 'structured', unavailable)
    def demo():
        pytest.fail('Live provider failure must not produce demo output')
    with pytest.raises(ProviderError, match='Model access denied'):
        agents.consult(Bus(), 'ProposalCreator', ProposalPlan, 'Plan response', {}, demo)
    assert messages[0][2] == 'error'


def test_provider_json_fence_preserves_schema_validation():
    from app.providers import validate_response
    from app.schemas import ProposalPlan
    from pydantic import ValidationError
    result = ProposalPlan(response_type='Technical response', formats=['docx'], outline=['Solution'], win_themes=['Security'], tone='Clear')
    assert validate_response(ProposalPlan, '```json\n' + result.model_dump_json() + '\n```') == result
    with pytest.raises(ValidationError):
        validate_response(ProposalPlan, '```json\n{"formats": ["invalid"]}\n```')
    with pytest.raises(ValidationError):
        validate_response(ProposalPlan, 'Some commentary\n```json\n' + result.model_dump_json() + '\n```')


def drain():
    while worker.run_once():
        pass


def setup(monkeypatch, client, title='RFP'):
    monkeypatch.setattr(worker, 'render_diagram', fake_render)
    monkeypatch.setattr(artifacts, 'render_diagram', fake_render)
    return client.post('/api/projects', json={'title': title}).json()['id']


def latest(client, pid):
    detail = client.get(f'/api/projects/{pid}').json()
    return detail, detail['proposals'][-1]


def test_orchestrator_routes_missing_information_to_user_then_drafts(workspace, monkeypatch):
    alice, _, sessions = workspace
    pid = setup(monkeypatch, alice)
    r = alice.post(f'/api/projects/{pid}/orchestrator', json={'message': VAGUE, 'formats': ['docx']})
    assert r.status_code == 202 and r.json()['intent'] == 'new_rfp'
    drain()
    _, p = latest(alice, pid)
    assert p['status'] == 'Awaiting Information'
    missing = {m['id'] for m in p['requirements']['missing'] if m['required']}
    assert {'timeline', 'commercial'} <= missing
    assert any(m['agent'] == 'Orchestrator' and m['recipient'] == 'RequirementsCollector' for m in p['messages'])
    # Free text while awaiting information is routed back to the RequirementsCollector.
    r = alice.post(f'/api/projects/{pid}/orchestrator', json={'message': 'Timeline is 3 months. Please proceed with your assumptions for the rest.'})
    assert r.json()['intent'] == 'information'
    drain()
    detail, p = latest(alice, pid)
    assert p['status'] == 'Pending Review', p['error']
    assert any(r['status'] == 'assumed' for r in p['requirements']['requirements'])
    agents = {m['agent'] for m in p['messages']}
    assert {'Orchestrator', 'RequirementsCollector', 'ProposalCreator', 'EstimationsAgent', 'ResourcesAgent', 'DeliveryMethodology', 'DocumentsAgent'} <= agents
    draft = next(v for v in detail['versions'] if v['id'] == p['draft_version_id'])
    kinds = {a['kind'] for a in draft['artifacts']}
    assert 'docx' in kinds and 'pptx' not in kinds and all(a['status'] == 'Completed' for a in draft['artifacts'])
    assert draft['response']['resources'] and draft['response']['methodology']['name']
    assert any(d['title'] == 'Delivery roadmap' for d in draft['response']['diagrams'])
    # Approval is never inferred from free text: it must be an explicit human decision.
    assert alice.post(f'/api/projects/{pid}/orchestrator', json={'message': 'Approved'}).status_code == 409


def test_review_changes_approve_publish_and_download(workspace, monkeypatch):
    alice, bob, sessions = workspace
    pid = setup(monkeypatch, alice)
    r = alice.post(f'/api/projects/{pid}/proposals', json={'request': COMPLETE, 'formats': ['docx', 'pptx']})
    assert r.status_code == 202
    prop_id = r.json()['proposal_id']
    assert alice.post(f'/api/projects/{pid}/proposals', json={'request': COMPLETE}).status_code == 409
    drain()
    _, p = latest(alice, pid)
    assert p['status'] == 'Pending Review', p['error']
    first_draft = p['draft_version_id']
    assert bob.post(f'/api/proposals/{prop_id}/review', json={'decision': 'approve'}).status_code == 404
    assert alice.post(f'/api/proposals/{prop_id}/review', json={'decision': 'request_changes', 'comments': 'no'}).status_code == 422
    r = alice.post(f'/api/proposals/{prop_id}/review', json={'decision': 'request_changes', 'comments': 'Add a stronger security section.'})
    assert r.status_code == 200
    drain()
    detail, p = latest(alice, pid)
    assert p['status'] == 'Pending Review' and p['draft_version_id'] != first_draft
    assert any('stronger security' in v['prompt'] or 'stronger security' in str(v['response']) for v in detail['versions']) or p['reviews'][0]['comments']
    r = alice.post(f'/api/proposals/{prop_id}/review', json={'decision': 'approve', 'comments': 'Good to submit.'})
    assert r.status_code == 200 and r.json()['documents'] >= 3
    detail, p = latest(alice, pid)
    assert p['status'] == 'Live' and p['published_at']
    docs = detail['response_documents']
    assert {d['kind'] for d in docs} >= {'docx', 'pptx', 'zip'}
    docx = next(d for d in docs if d['kind'] == 'docx')
    assert bob.post(f'/api/response-documents/{docx["id"]}/download-link').status_code == 404
    url = alice.post(f'/api/response-documents/{docx["id"]}/download-link').json()['url']
    assert bob.get(url).status_code == 404
    data = alice.get(url).content
    text = '\n'.join(par.text for par in Document(io.BytesIO(data)).paragraphs)
    assert 'Our understanding of your requirements' in text and 'Delivery methodology' in text
    pptx = next(d for d in docs if d['kind'] == 'pptx')
    deck = Presentation(io.BytesIO(alice.get(alice.post(f'/api/response-documents/{pptx["id"]}/download-link').json()['url']).content))
    titles = ' '.join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert 'Proposed team' in titles and 'RFP proposal response' in titles
    assert alice.post(f'/api/proposals/{prop_id}/review', json={'decision': 'approve'}).status_code == 409
    # Project deletion removes published response documents from storage.
    with sessions() as db:
        keys = [d.key for d in db.query(ResponseDocument).all()]
    assert alice.delete(f'/api/projects/{pid}').status_code == 200
    from app.storage import storage
    assert all(not (storage.root / k).exists() for k in keys)


def test_approval_blocked_when_files_failed_and_failed_proposal_retry(workspace, monkeypatch):
    alice, _, sessions = workspace
    pid = setup(monkeypatch, alice)
    def broken(source):
        raise RuntimeError('renderer down')
    monkeypatch.setattr(artifacts, 'render_diagram', broken)
    prop_id = alice.post(f'/api/projects/{pid}/proposals', json={'request': COMPLETE, 'formats': ['docx']}).json()['proposal_id']
    drain()
    _, p = latest(alice, pid)
    assert p['status'] == 'Pending Review'
    r = alice.post(f'/api/proposals/{prop_id}/review', json={'decision': 'approve'})
    assert r.status_code == 409 and 'Retry failed files' in r.json()['detail']
    assert alice.post(f'/api/proposals/{prop_id}/retry').status_code == 409
    with sessions() as db:
        proposal = db.get(Proposal, prop_id)
        proposal.status = 'Failed'
        db.commit()
    r = alice.post(f'/api/projects/{pid}/orchestrator', json={'message': 'retry'})
    assert r.status_code == 202 and r.json()['intent'] == 'retry'
    monkeypatch.setattr(artifacts, 'render_diagram', fake_render)
    drain()
    _, p = latest(alice, pid)
    assert p['status'] == 'Pending Review'
    with sessions() as db:
        assert db.query(AgentMessage).filter_by(proposal_id=prop_id, agent='Orchestrator').count() >= 2
        assert db.query(Artifact).filter_by(version_id=p['draft_version_id'], status='Failed').count() == 0
