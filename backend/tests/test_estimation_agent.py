import io
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import ValidationError
from app import agents
from app.config import settings
from app.schemas import Estimate, PricingInputs, EstimationResult, RequirementsAssessment, calculate
from app.artifacts import create_document
from app.diagrams import DiagramsAgent, delivery_views
from test_workspace import sample_solution, fake_render
from test_presentation import visible_text


def scope(model='fixed_price'):
    s = sample_solution()
    s.estimates = [Estimate(activity='Build portal',role='Engineer',low=8,expected=10,high=12,rate=100,confidence='medium')]
    s.contingency_percent = 20
    s.commercial_model = model
    s.pricing = PricingInputs(variations=[{'name':'Add reporting','role':'Engineer','effort_days':2}])
    return s


def test_fixed_price_and_tm_are_distinct_and_variations_are_separate():
    s = scope()
    fixed = calculate(s)['commercial']
    assert fixed['base_cost'] == {'low':800,'expected':1000,'high':1200}
    assert fixed['indicative_fixed_price'] == 1200
    assert fixed['variation_cost'] == 200 and fixed['expected_with_variations'] == 1400
    s.commercial_model = 'capped_time_and_materials'
    s.pricing.budget_cap = 1100
    tm = calculate(s)['commercial']
    assert tm['indicative_fixed_price'] is None
    assert tm['expected_with_variations'] == 1200
    assert tm['exceeds_cap'] is True  # Never hide effort cost behind the cap.


def test_missing_rates_zero_rates_and_unknown_variation_roles():
    s = scope('time_and_materials')
    s.estimates[0].rate = None
    assert calculate(s)['commercial']['expected_with_variations'] is None
    assert calculate(s)['commercial']['variation_cost'] is None
    s.estimates[0].rate = 0
    assert calculate(s)['commercial']['expected_with_variations'] == 0
    s.pricing.variations[0].role = 'Unpriced specialist'
    assert calculate(s)['commercial']['variation_cost'] is None
    s.pricing.variations[0].rate = 75.25
    assert calculate(s)['commercial']['variation_cost'] == 150.5


def test_negative_changes_and_ambiguous_role_rates_require_review():
    s = scope('time_and_materials')
    s.pricing.variations[0].effort_days = -2
    assert calculate(s)['commercial']['expected_with_variations'] == 800
    s.pricing.variations[0].effort_days = -20
    result = calculate(s)['commercial']
    assert result['expected_with_variations'] is None and any('reduction' in value for value in result['missing_inputs'])
    s = scope()
    s.estimates.append(s.estimates[0].model_copy(update={'activity':'Support','rate':120}))
    assert calculate(s)['commercial']['variation_cost'] is None


def test_other_contracts_flag_missing_billing_terms_and_decimal_math():
    s = scope('retainer')
    assert any('billing schedule' in value for value in calculate(s)['commercial']['missing_inputs'])
    s.estimates[0].expected = 10.1
    s.estimates[0].rate = 100.15
    assert calculate(s)['commercial']['base_cost']['expected'] == 1011.52
    with pytest.raises(ValidationError):
        PricingInputs(role_rates=[{'role':'Engineer','rate':100},{'role':'engineer','rate':200}])


def test_solution_delegates_scope_to_estimator_and_estimator_calls_diagrams(monkeypatch):
    monkeypatch.setattr(settings, 'llm_provider', 'bedrock')
    s = scope()
    messages = []
    class Bus:
        def send(self, *message): messages.append(message)
    request = {'current_prompt':'Fixed price portal with reporting change',
               'preferences':{'outputs':['estimates','presentation'],'pricing':{'model':'fixed_price','currency':'USD','role_rates':[{'role':'Engineer','rate':100}],
                   'variations':[{'name':'Add reporting','role':'Engineer','effort_days':2}]}}}
    def structured(model, system, payload, validate):
        assert model is EstimationResult
        assert payload['draft_estimates'][0]['expected'] == 10
        assert payload['scope_detail'] and payload['pricing']['model'] == 'fixed_price'
        # A model's invented rate must never change the supplied price basis.
        result = EstimationResult(estimates=[s.estimates[0].model_copy(update={'rate':999999})],
                                  contingency_percent=20,commercial_model='time_and_materials',commercial_basis='Fixed scope pending approval.')
        return validate(result)
    monkeypatch.setattr(agents, 'structured', structured)
    phases = []
    result = agents.SolutionAgent(Bus()).run(request,before_estimates=lambda:phases.append('Estimating'),generate_fn=lambda _:s)
    assert phases == ['Estimating']
    assert result.commercial_model == 'fixed_price' and result.estimates[0].rate == 100
    assert calculate(result)['commercial']['expected_with_variations'] == 1400
    assert any(m[:3] == ('SolutionAgent','EstimationsAgent','task') for m in messages)
    assert any(m[:3] == ('EstimationsAgent','Diagrams Agent','task') for m in messages)
    assert any(m[:3] == ('Diagrams Agent','EstimationsAgent','result') for m in messages)
    assert any(m[:3] == ('EstimationsAgent','SolutionAgent','result') for m in messages)
    assert any(d.kind == 'estimations' and d.title == 'Estimations Summary' for d in result.diagrams)


def test_estimation_summary_slide_uses_calculated_values_and_revisions():
    s = scope()
    d = DiagramsAgent().estimation_summary(s)
    assert '1,400.00' in d.source and 'fixed' in d.description.lower()
    s.diagrams = [d]
    data = create_document('pptx',s,'demo',[(d.title,fake_render(d.source)['png'])])
    deck = Presentation(io.BytesIO(data))
    slide = next(slide for slide in deck.slides if any(sh.has_text_frame and sh.text == 'Estimations Summary' for sh in slide.shapes))
    assert sum(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes) == 1
    assert 'Commercial estimation basis' in visible_text(deck) and 'Proposed estimate variations' in visible_text(deck)
    s.estimates[0].rate = 200
    refreshed,_ = delivery_views(s)
    assert '2,800.00' in next(d for d in refreshed if d.kind == 'estimations').source
