"""Bounded Office template validation and private template selection."""
import io
import zipfile
from pathlib import PurePosixPath
from lxml import etree
from docx import Document
from pptx import Presentation
from fastapi import HTTPException
from .db import ProposalTemplate, Project
from .storage import storage
from .template_formats import extract_format

TYPES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',
    'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',
}


def validate_template(filename, data):
    ext = filename.rsplit('.', 1)[-1].lower()
    kind = {'docx': 'docx', 'dotx': 'docx', 'pptx': 'pptx', 'potx': 'pptx'}.get(ext)
    if not kind:
        raise ValueError('Upload a Word (.docx, .dotx) or PowerPoint (.pptx, .potx) template.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            entries = package.infolist()
            if len(entries) > 3000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                raise ValueError('Template expands beyond the supported size limit.')
            for entry in entries:
                path = PurePosixPath(entry.filename)
                name = entry.filename.lower()
                if path.is_absolute() or '..' in path.parts or '\\' in name or entry.flag_bits & 1:
                    raise ValueError('Invalid or encrypted template package.')
                if any(part in name for part in ('vbaproject', '/embeddings/', '/activex/')):
                    raise ValueError('Templates with macros, embedded objects or ActiveX are not supported.')
            if len({e.filename for e in entries}) != len(entries):
                raise ValueError('Template contains duplicate package entries.')
            # Normalize native template content types for the document libraries.
            root = etree.fromstring(package.read('[Content_Types].xml'), etree.XMLParser(resolve_entities=False, no_network=True))
            main = '/word/document.xml' if kind == 'docx' else '/ppt/presentation.xml'
            found = False
            for item in root:
                if item.get('PartName') == main:
                    allowed = [TYPES[kind], TYPES[kind].replace('document.main', 'template.main') if kind == 'docx' else TYPES[kind].replace('presentation.main', 'template.main')]
                    if item.get('ContentType') not in allowed:
                        raise ValueError('Template type does not match the file extension.')
                    item.set('ContentType', TYPES[kind])
                    found = True
            if not found:
                raise ValueError('Template is missing its main document part.')
            out = io.BytesIO()
            with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as normalized:
                for entry in entries:
                    normalized.writestr(entry.filename, etree.tostring(root) if entry.filename == '[Content_Types].xml' else package.read(entry))
        normalized_data = out.getvalue()
        doc = Document(io.BytesIO(normalized_data)) if kind == 'docx' else Presentation(io.BytesIO(normalized_data))
        if kind == 'pptx' and (not doc.slide_layouts or len(doc.slides) > 100):
            raise ValueError('PowerPoint templates need a slide layout and at most 100 sample slides.')
        return kind, normalized_data
    except ValueError:
        raise
    except Exception:
        raise ValueError('This file is not a readable, unencrypted Office template.') from None


def validate_selection(db, user, preferences):
    for field, kind in [('word_template_id', 'docx'), ('powerpoint_template_id', 'pptx')]:
        template_id = getattr(preferences, field)
        if template_id:
            row = db.get(ProposalTemplate, template_id)
            if not row or row.owner_id != user.id or row.archived or row.kind != kind:
                raise HTTPException(400, 'The selected template is unavailable. Choose a template from your library.')


def template_for_version(db, version, kind):
    field = {'docx': 'word_template_id', 'pptx': 'powerpoint_template_id'}.get(kind)
    template_id = version.preferences.get(field) if field else None
    if not template_id:
        return None
    row = db.get(ProposalTemplate, template_id)
    project = db.get(Project, version.project_id)
    if not row or row.owner_id != project.owner_id or row.kind != kind:
        raise ValueError('The selected proposal template is unavailable.')
    # Archived templates remain available to existing versions and their retries.
    return prepare_template(row.filename, storage.get(row.key))[1]


def prepare_template(filename, data):
    kind, normalized = validate_template(filename, data)
    try:
        return kind, extract_format(kind, normalized)
    except Exception:
        raise ValueError('Could not extract a reusable format. Open and save this file in Word or PowerPoint, then retry.') from None
