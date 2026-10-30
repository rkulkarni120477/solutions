"""Extract reusable Office formatting without retaining the uploaded content."""
import io
from collections import Counter
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from pptx import Presentation
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'

# Only formatting parts may remain reachable in the saved Office package.
FORMAT_RELATIONS = {
    'styles', 'styleswitheffects', 'numbering', 'fonttable', 'font', 'settings',
    'websettings', 'theme', 'header', 'footer', 'slidemaster', 'slidelayout',
    'presprops', 'viewprops', 'tablestyles',
}


def prune_content_parts(package):
    parts = list(package.iter_parts())
    root_rels = package.rels if hasattr(package, 'rels') else package._rels
    for rid, rel in list(root_rels.items()):
        if rel.is_external or not rel.reltype.endswith('/officeDocument'):
            root_rels.pop(rid)
    for part in parts:
        for rid, rel in list(part.rels.items()):
            if rel.is_external or rel.reltype.rsplit('/', 1)[-1].lower() not in FORMAT_RELATIONS:
                part.rels.pop(rid)
        # Theme parts can also carry text or image references in extensions.
        if str(part.partname).startswith(('/word/theme/', '/ppt/theme/')):
            root = etree.fromstring(part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
            clear_text_and_content(root)
            part._blob = etree.tostring(root)


def clear_text_and_content(element):
    remove = {
        (W, 'drawing'), (W, 'pict'), (W, 'object'), (W, 'hyperlink'),
        (W, 'instrText'), (W, 'fldChar'), (W, 'docVars'), (W, 'mailMerge'),
        (W, 'attachedTemplate'), (W, 'bookmarkStart'), (W, 'bookmarkEnd'),
        (W, 'sdtPr'), (W, 'smartTagPr'), (W, 'customXmlPr'),
        (W, 'commentRangeStart'), (W, 'commentRangeEnd'), (W, 'commentReference'),
        (W, 'footnoteReference'), (W, 'endnoteReference'), (W, 'printerSettings'),
        (W, 'ins'), (W, 'del'), (W, 'moveFrom'), (W, 'moveTo'),
        (P, 'pic'), (P, 'graphicFrame'), (P, 'oleObj'), (P, 'extLst'),
        (A, 'hlinkClick'), (A, 'hlinkMouseOver'), (A, 'blipFill'), (A, 'extLst'),
    }
    for node in list(element.iter()):
        if not isinstance(node.tag, str):
            continue
        name = etree.QName(node)
        if (name.namespace, name.localname) in remove:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
        elif (name.namespace, name.localname) in {(W, 't'), (W, 'delText'), (A, 't')}:
            node.text = ''
        if name.localname == 'cNvPr':
            node.attrib.pop('descr', None)
            node.attrib.pop('title', None)
            node.set('name', 'Format shape ' + node.get('id', ''))


def word_format(data):
    doc = Document(io.BytesIO(data))
    # Capture direct body typography as well as named styles. Many real-world
    # samples format runs directly instead of changing the Normal style.
    body_paragraphs = [p for p in doc.paragraphs[:1000] if p.text.strip() and not p.style.name.startswith(('Heading', 'Title', 'Subtitle'))]
    if body_paragraphs:
        paragraph = max(body_paragraphs, key=lambda p: len(p.text))
        runs = [r for r in paragraph.runs if r.text.strip() and r._r.rPr is not None]
        if runs:
            source = max(runs, key=lambda r: len(r.text))._r.rPr
            target = doc.styles['Normal'].element.get_or_add_rPr()
            for tag in ['rFonts', 'sz', 'szCs', 'color']:
                value = source.find(f'{{{W}}}{tag}')
                if value is not None:
                    old = target.find(f'{{{W}}}{tag}')
                    if old is not None:
                        target.remove(old)
                    target.append(deepcopy(value))
        for prop in ['alignment', 'space_before', 'space_after', 'line_spacing', 'left_indent', 'right_indent', 'first_line_indent']:
            value = getattr(paragraph.paragraph_format, prop)
            if value is not None:
                setattr(doc.styles['Normal'].paragraph_format, prop, value)
    section = deepcopy(doc.sections[0]._sectPr)
    body = doc.element.body
    for child in list(body):
        body.remove(child)
    body.append(OxmlElement('w:p'))
    body.append(section)
    # Headers/footers retain their layout and paragraph formatting, not their text.
    for part in list(doc.part.package.iter_parts()):
        element = getattr(part, 'element', None)
        if element is not None and str(part.partname).startswith(('/word/header', '/word/footer', '/word/settings')):
            clear_text_and_content(element)
    prune_content_parts(doc.part.package)
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


def powerpoint_format(data):
    prs = Presentation(io.BytesIO(data))
    fonts = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += len(run.text)
    # A sample may apply its background directly on slides rather than layouts.
    if prs.slides:
        background = prs.slides[0]._element.xpath('./p:cSld/p:bg')
        layout = min(prs.slide_layouts, key=lambda item: len(item.placeholders))
        if background and not layout._element.xpath('./p:cSld/p:bg'):
            layout._element.cSld.insert(0, deepcopy(background[0]))
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    for tag in ['custShowLst', 'notesMasterIdLst', 'handoutMasterIdLst']:
        for node in prs._element.findall(f'{{{P}}}{tag}'):
            prs._element.remove(node)
    for master in prs.slide_masters:
        clear_text_and_content(master._element)
        for layout in master.slide_layouts:
            clear_text_and_content(layout._element)
        if fonts:
            for rel in master.part.rels.values():
                if rel.reltype.endswith('/theme'):
                    theme = etree.fromstring(rel.target_part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
                    for node in theme.findall(f'.//{{{A}}}minorFont/{{{A}}}latin'):
                        node.set('typeface', fonts.most_common(1)[0][0])
                    rel.target_part._blob = etree.tostring(theme)
    clear_text_and_content(prs._element)
    prune_content_parts(prs.part.package)
    result = io.BytesIO()
    prs.save(result)
    return result.getvalue()


def extract_format(kind, data):
    return word_format(data) if kind == 'docx' else powerpoint_format(data)
