"""Editable, content-aware project briefing slides generated from a validated solution."""
import io
import json
import re
import textwrap
from decimal import Decimal

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .schemas import Solution, calculate

NAVY = '14273F'
BLUE = '315EA5'
TEAL = '327F7A'
INK = '263D58'
MUTED = '62758B'
PALE = 'EFF4FA'
WHITE = 'FFFFFF'
WIDTH, HEIGHT = 13.333, 7.5


def plain(text):
    text = re.sub(r'!\[([^]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'(?m)^\s*#{1,6}\s*', '', text)
    text = re.sub(r'(?m)^\s*[-*+]\s+', '', text)
    text = re.sub(r'\*\*|__|`', '', text)
    return text.strip()


def segments(text):
    """Preserve words and sentences, paginate by wrapped lines, and omit raw code/JSON."""
    text = re.sub(r'```[\s\S]*?```', 'Code details are available in the solution document.', text)
    for paragraph in re.split(r'\n\s*\n', text):
        paragraph = plain(paragraph)
        if not paragraph:
            continue
        try:
            if isinstance(json.loads(paragraph), (dict, list)):
                paragraph = 'Structured configuration details are available in the solution document.'
        except (ValueError, TypeError):
            pass
        group = ''
        for sentence in re.split(r'(?<=[.!?])\s+|\n', paragraph):
            if group and len(group) + len(sentence) + 1 > 380:
                yield group
                group = ''
            if len(sentence) > 380:
                if group:
                    yield group
                    group = ''
                yield from textwrap.wrap(sentence, width=380, break_long_words=True, break_on_hyphens=False)
            elif sentence:
                group = (group + ' ' + sentence).strip()
        if group:
            yield group


def pages(items, width=95, max_lines=16):
    page, used = [], 0
    for item in items:
        item = plain(str(item))
        if not item:
            continue
        # Split oversized entries into continuation paragraphs before packing a page.
        for part in textwrap.wrap(item, width=width * (max_lines - 2), break_long_words=True, break_on_hyphens=False):
            lines = len(textwrap.wrap(part, width=width, break_long_words=True)) + 1
            if page and used + lines > max_lines:
                yield page
                page, used = [], 0
            page.append(part)
            used += lines
    if page:
        yield page


class Briefing:
    def __init__(self, solution, provider, metadata=None):
        self.solution = solution
        self.provider = provider
        self.metadata = metadata or {}
        self.prs = Presentation()
        self.prs.slide_width = Inches(WIDTH)
        self.prs.slide_height = Inches(HEIGHT)
        self.prs.core_properties.title = solution.title
        self.prs.core_properties.subject = 'Project solution briefing and estimates'
        self.prs.core_properties.author = 'Solution Architect Workspace'
        self.label = 'DEMO: illustrative output. No live LLM response.' if provider == 'demo' else f'Provider: {provider}. Architect review required.'

    def text(self, slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(str(text).split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Calibri'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.alignment = align
            p.space_after = Pt(9)
            p.line_spacing = 1.12
        return box

    def rect(self, slide, x, y, w, h, fill):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        return shape

    def slide(self, title, subtitle='', dark=False):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY if dark else WHITE)
        display_title = textwrap.shorten(title, width=110, placeholder='…')
        self.text(slide, display_title, .65, .65, 12, 1, 30 if len(display_title) < 45 else 24 if len(display_title) < 85 else 20, WHITE if dark else NAVY, True)
        if subtitle:
            self.text(slide, subtitle, .68, 1.62, 11.95, .48, 11, 'BDCEE3' if dark else MUTED)
        self.text(slide, self.label, .68, 7.04, 10.9, .23, 9, 'BDCEE3' if dark else MUTED)
        self.text(slide, f'{len(self.prs.slides):02d}', 12.05, 7.02, .55, .3, 10, 'BDCEE3' if dark else MUTED, align=PP_ALIGN.RIGHT)
        return slide

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = str(text)

    def narrative(self, title, text, subtitle='', sources=None):
        chunks = list(pages(segments(text)))
        for index, chunk in enumerate(chunks):
            heading = title if index == 0 else f'{title} (continued {index + 1})'
            slide = self.slide(heading, subtitle)
            self.text(slide, '\n'.join(chunk), .7, 2.25, 11.9, 4.35, 18)
            self.notes(slide, title + '\n\n' + text + ('\n\nReference passage IDs:\n' + '\n'.join(sources) if sources else ''))
        return chunks

    def table(self, title, headers, rows, widths, subtitle='', per_page=6, note=''):
        def row_height(values):
            lines = max(len(textwrap.wrap(str(value), max(8, int((widths[c] - .2) * 9)))) or 1 for c, value in enumerate(values))
            return max(.5, .15 + lines * .21)
        groups, group, used = [], [], row_height(headers)
        for row in rows:
            height = row_height(row)
            if group and (used + height > 4.05 or len(group) >= per_page):
                groups.append(group)
                group, used = [], row_height(headers)
            group.append(row)
            used += height
        if group:
            groups.append(group)
        for index, subset in enumerate(groups):
            suffix = f' ({index + 1})' if len(groups) > 1 else ''
            slide = self.slide(title + suffix, subtitle)
            table = slide.shapes.add_table(len(subset) + 1, len(headers), Inches(.68), Inches(2.2), Inches(11.95), Inches(.5 + .57 * len(subset))).table
            for i, width in enumerate(widths):
                table.columns[i].width = Inches(width)
            for r, values in enumerate([headers, *subset]):
                row_lines = 1
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    cell.text = plain(str(value))
                    cell.margin_left = cell.margin_right = Inches(.1)
                    cell.margin_top = cell.margin_bottom = Inches(.08)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else WHITE if r % 2 == 0 else PALE)
                    for p in cell.text_frame.paragraphs:
                        p.font.name = 'Calibri'
                        p.font.size = Pt(12 if r == 0 else 13)
                        p.font.bold = r == 0
                        p.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else INK)
                        p.space_after = Pt(0)
                    capacity = max(8, int((widths[c] - .2) * 9))
                    row_lines = max(row_lines, len(textwrap.wrap(cell.text, capacity)))
                table.rows[r].height = Inches(max(.5, .15 + row_lines * .21))
            if note:
                self.text(slide, note, .7, 6.55, 11.9, .35, 10, MUTED)
            self.notes(slide, '\n'.join(' / '.join(map(str, row)) for row in [headers, *subset]))


def number(value):
    return f'{value:,.2f}'.rstrip('0').rstrip('.')


def create_presentation(solution: Solution, provider: str, diagrams=None, metadata=None):
    deck = Briefing(solution, provider, metadata)
    s, totals = solution, calculate(solution)
    proposal = bool(deck.metadata.get('proposal'))
    cover = deck.slide('RFP proposal response' if proposal else 'Project solution briefing', dark=True)
    deck.text(cover, s.title, .72, 2.1, 11.9, 2.45, 32 if len(s.title) < 95 else 28, WHITE, True)
    deck.text(cover, 'Solution, estimates, team, methodology and delivery' if proposal else 'Executive summary, solution design, estimates and delivery', .74, 4.75, 11.8, .6, 18, 'BDCEE3')
    info = []
    if deck.metadata.get('version') is not None:
        info.append(f"Response version {deck.metadata['version']}")
    if deck.metadata.get('created_at'):
        info.append(str(deck.metadata['created_at'])[:10])
    deck.text(cover, '   '.join(info), .74, 5.85, 11.8, .4, 12, 'BDCEE3')
    deck.notes(cover, s.title + '\n' + deck.label)

    deck.narrative('Executive summary', s.summary, 'Project objectives and proposed approach')
    # Put the proposed design first, then preserve other solution sections in their original order.
    proposed = [section for section in s.sections if re.search(r'propos|solution|approach|architecture', section.title, re.I)]
    remaining = [section for section in s.sections if section not in proposed]
    if not proposed:
        deck.narrative('Proposed solution', 'A separate solution narrative was not provided. Review the executive summary and detailed response before making architecture decisions.')
    for section in proposed + remaining:
        # Parse actual Markdown tables into editable PowerPoint tables.
        lines = section.content.splitlines()
        text, tables = [], []
        i = 0
        while i < len(lines):
            if i + 1 < len(lines) and '|' in lines[i] and re.match(r'^\s*\|?\s*:?-{3,}', lines[i + 1]):
                headers = [plain(c) for c in lines[i].strip().strip('|').split('|')]
                rows = []
                i += 2
                while i < len(lines) and '|' in lines[i] and lines[i].strip():
                    values = [plain(c) for c in lines[i].strip().strip('|').split('|')]
                    if len(values) == len(headers):
                        rows.append(values)
                    i += 1
                if rows and 1 <= len(headers) <= 6:
                    tables.append((headers, rows))
                else:
                    text.extend(' / '.join(row) for row in [headers, *rows])
            else:
                text.append(lines[i])
                i += 1
        deck.narrative(section.title, '\n'.join(text), section.classification.capitalize(), section.sources)
        for headers, rows in tables:
            # Expand very long cells into continuation rows so no source text is clipped.
            expanded = []
            limit = max(40, int(210 / len(headers)))
            for row in rows:
                parts = [textwrap.wrap(value, limit, break_long_words=True) or [''] for value in row]
                for n in range(max(map(len, parts))):
                    expanded.append([part[n] if n < len(part) else '' for part in parts])
            deck.table(section.title, headers, expanded, [11.95 / len(headers)] * len(headers), section.classification.capitalize(), per_page=7)

    rendered = {title: png for title, png in diagrams or []}
    for diagram in s.diagrams:
        slide = deck.slide(diagram.title, diagram.kind.capitalize() + ' architecture. Editable Mermaid source is included separately in the package.')
        png = rendered.get(diagram.title)
        if png:
            image = Image.open(io.BytesIO(png))
            ratio = min(11.85 / image.width, 4.55 / image.height)
            w, h = image.width * ratio, image.height * ratio
            deck.rect(slide, .68, 2.15, 11.95, 4.65, PALE)
            slide.shapes.add_picture(io.BytesIO(png), Inches((WIDTH - w) / 2), Inches(2.2 + (4.55 - h) / 2), width=Inches(w), height=Inches(h))
        else:
            deck.text(slide, 'Diagram export is unavailable. Retry the failed diagram in Files, then refresh the presentation. The Mermaid source remains available for review.', .75, 2.7, 11.7, 2.2, 21)
        deck.notes(slide, diagram.source)
    if not s.diagrams:
        deck.narrative('Architecture', 'No diagram was requested or provided for this response. Enable Diagrams when generating a new version to include rendered architecture views.')

    slide = deck.slide('Estimation overview', 'Indicative effort in person-days. Contingency included. Implementation and operating costs remain separate.')
    if s.estimates:
        for i, (key, label) in enumerate([('low', 'LOW'), ('expected', 'EXPECTED'), ('high', 'HIGH')]):
            x = .75 + i * 4.15
            deck.text(slide, label, x, 2.5, 3.65, .35, 12, MUTED, True)
            deck.text(slide, number(totals['with_contingency'][key]), x, 3.05, 3.65, .9, 46, BLUE if key == 'expected' else NAVY, True)
            deck.text(slide, 'person-days', x, 4.02, 3.65, .4, 16, MUTED)
        deck.text(slide, f'Contingency: {number(s.contingency_percent)}%     Staffing assumption: {number(s.staffing_fte)} FTE', .75, 5.2, 11.8, .5, 18, INK)
        deck.text(slide, f"Capacity lower bound: {number(totals['capacity_floor_weeks'])} working weeks at 5 days/week. Dependencies and role availability can extend elapsed delivery.", .75, 5.95, 11.8, .7, 14, MUTED)
    else:
        deck.text(slide, 'No effort estimate was requested or provided. Select Estimates for a new response to populate workstreams, roles and calculated totals.', .75, 2.8, 11.75, 2, 22)
    deck.notes(slide, json.dumps(totals, indent=2))

    if s.estimates:
        rows = []
        basis = []
        rates = []
        for row in s.estimates:
            activity = textwrap.wrap(row.activity, 75) or ['']
            role = textwrap.wrap(row.role, 45) or ['']
            for n in range(max(len(activity), len(role))):
                rows.append([activity[n] if n < len(activity) else '', role[n] if n < len(role) else '', *([number(row.low), number(row.expected), number(row.high)] if n == 0 else ['', '', ''])])
            basis.append([row.activity, row.role, row.dependency or 'Not specified', row.confidence.capitalize()])
            rates.append(f'{row.activity}: ' + (f'{s.currency} {number(row.rate)} per person-day' if row.rate is not None else 'daily rate not supplied'))
        deck.table('Effort by workstream', ['Activity', 'Role', 'Low', 'Expected', 'High'], rows, [4.25, 3.2, 1.5, 1.5, 1.5], 'Base effort in person-days, before contingency', per_page=5,
                   note=f"Base totals: {number(totals['effort']['low'])} low / {number(totals['effort']['expected'])} expected / {number(totals['effort']['high'])} high person-days")
        # Native editable chart data, paginated for large estimate sets.
        for start in range(0, len(s.estimates), 8):
            group = s.estimates[start:start + 8]
            slide = deck.slide('Expected effort by workstream' + (f' ({start // 8 + 1})' if len(s.estimates) > 8 else ''), 'Base person-days, before contingency. Chart values remain editable in PowerPoint.')
            data = CategoryChartData()
            data.categories = [textwrap.shorten(row.activity, width=42, placeholder='…') for row in group]
            data.add_series('Expected person-days', [row.expected for row in group])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(.75), Inches(2.25), Inches(11.8), Inches(4.3), data).chart
            chart.has_legend = False
            chart.category_axis.tick_labels.font.size = Pt(12)
            chart.value_axis.tick_labels.font.size = Pt(11)
            chart.value_axis.minimum_scale = 0
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.size = Pt(12)
            plot.series[0].format.fill.solid()
            plot.series[0].format.fill.fore_color.rgb = RGBColor.from_string(BLUE)
            deck.notes(slide, '\n'.join(f'{row.activity}: {row.expected} person-days' for row in group))
        deck.table('Estimate confidence and dependencies', ['Activity', 'Role', 'Dependencies', 'Confidence'], basis, [3.4, 2.5, 4.55, 1.5], 'Staffing, parallel work and external dependencies require validation', per_page=5)

    costs = totals['implementation_cost']
    cost_text = ('\n\n'.join(f'{label}: {s.currency} {number(costs[key])}' for key, label in [('low', 'Low implementation cost'), ('expected', 'Expected implementation cost'), ('high', 'High implementation cost')]) + f'\n\nIncludes {number(s.contingency_percent)}% contingency. Rates are per person-day.' if costs else 'Implementation cost is not priced. Supply a daily rate for every workstream to calculate low, expected and high implementation costs.')
    deck.narrative('Implementation cost', cost_text, 'Calculated from supplied rates only')
    if s.estimates:
        deck.narrative('Rate assumptions', '\n\n'.join(rates), f'Currency: {s.currency}. Daily rates are distinct from total cost.')
    deck.narrative('Operating costs and exclusions', 'Recurring operating costs\n\n' + (s.operating_cost_note or 'No operating-cost information was supplied.') + '\n\nExclusions\n\n' + ('\n\n'.join(s.exclusions) or 'No exclusions were provided.'), 'Recurring consumption and licenses are separate from implementation effort')

    if s.phases:
        for start in range(0, len(s.phases), 4):
            group = s.phases[start:start + 4]
            slide = deck.slide('Delivery roadmap' + (f' ({start // 4 + 1})' if len(s.phases) > 4 else ''), 'Sequence, durations and dependencies from the project response')
            for i, phase in enumerate(group):
                x = .75 + i * 3.1
                circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(2.3), Inches(.62), Inches(.62))
                circle.fill.solid()
                circle.fill.fore_color.rgb = RGBColor.from_string(BLUE)
                circle.line.fill.background()
                deck.text(slide, f'{start + i + 1:02d}', x, 2.46, .62, .35, 14, WHITE, True, PP_ALIGN.CENTER)
                deck.text(slide, textwrap.shorten(phase.name, width=44, placeholder='…'), x, 3.13, 2.65, 1.3, 18, NAVY, True)
                deck.text(slide, textwrap.shorten(phase.duration, width=45, placeholder='…'), x, 4.65, 2.65, .75, 14, TEAL)
                deck.text(slide, f'{len(phase.activities)} defined activities', x, 5.7, 2.65, .45, 12, MUTED)
            deck.notes(slide, '\n\n'.join(f'{p.name}\n{p.duration}\n' + '\n'.join(p.activities) + '\nDependencies: ' + p.dependencies for p in group))
        for phase in s.phases:
            deck.narrative(phase.name, '\n\n'.join(phase.activities) + '\n\nDependencies: ' + phase.dependencies, phase.duration)
        deck.narrative('Delivery duration and staffing', s.duration_note or 'Delivery duration requires validation.', f'Staffing assumption: {number(s.staffing_fte)} FTE. Effort and elapsed duration are different measures.')
    else:
        deck.narrative('Delivery plan', 'No implementation plan was requested or provided. Select Implementation plan for a new response to populate delivery phases and dependencies.')
    if s.methodology:
        m = s.methodology
        deck.narrative(f'Delivery methodology: {m.name}', m.overview + ''.join(f'\n\n{label}\n\n' + '\n\n'.join(items) for label, items in [('Ceremonies', m.ceremonies), ('Governance', m.governance), ('Quality gates', m.quality_gates)] if items), 'How the engagement will be delivered and governed')
    if s.resources:
        deck.table('Proposed team', ['Role', 'Count', 'Allocation', 'Responsibilities'], [[r.role, number(r.count), r.allocation, r.responsibilities] for r in s.resources], [2.8, 1.1, 2.7, 5.35], f'Team sized from the estimate. Staffing assumption: {number(s.staffing_fte)} FTE.', per_page=5)
    deck.narrative('Assumptions for review', '\n\n'.join(s.assumptions) or 'No assumptions were recorded in the response.', 'Confirm assumptions before approving scope or estimates')
    deck.narrative('Open questions and decisions', '\n\n'.join(s.questions) or 'No unresolved questions were recorded. Architect approval is still required.', 'Outstanding inputs for project stakeholders')
    closing = deck.slide('Project review and next steps', dark=True)
    deck.text(closing, 'Architecture review\nScope and estimate validation\nStakeholder decisions\nDelivery approval', .8, 2.4, 11.6, 3.5, 27, WHITE)
    deck.notes(closing, 'Proposed review workflow. Approval status is not implied.\n' + deck.label)
    buffer = io.BytesIO()
    deck.prs.save(buffer)
    return buffer.getvalue()
