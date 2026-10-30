import io
import re
import zipfile
import subprocess
import sys
import json
import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from .config import settings


def extract_isolated(filename, data):
    try:
        env = {**os.environ, 'PYTHONPATH': str(Path(__file__).resolve().parents[1]) + os.pathsep + os.environ.get('PYTHONPATH', '')}
        result = subprocess.run([sys.executable, '-m', 'app.extract_task', filename], input=data,
                                capture_output=True, timeout=45, check=False, env=env)
        content = json.loads(result.stdout)
        if result.returncode or 'error' in content:
            raise ValueError(content.get('error', 'Extraction failed.'))
        return content['passages'], content['warning']
    except (subprocess.TimeoutExpired, json.JSONDecodeError):
        raise ValueError('Document processing exceeded its limits. Split the document or upload a text-based copy.') from None


def extract(filename: str, data: bytes):
    ext = Path(filename).suffix.lower()
    passages = []
    ocr_truncated = False
    if ext in {'.txt', '.md'}:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('Text documents must use UTF-8 encoding.') from exc
        if '\x00' in text:
            raise ValueError('Binary content is not a text document.')
        passages = [{'reference': f'section {i + 1}', 'text': part} for i, part in enumerate(re.split(r'\n(?=#{1,6} )', text))]
    elif ext == '.pdf':
        if not data.startswith(b'%PDF-'):
            raise ValueError('The file does not contain a valid PDF signature.')
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('Encrypted PDFs are not supported. Upload a decrypted copy.')
        if len(reader.pages) > 250:
            raise ValueError('PDFs are limited to 250 pages. Split this document first.')
        passages = [{'reference': f'page {i + 1}', 'text': page.extract_text() or ''} for i, page in enumerate(reader.pages)]
        if not any(p['text'].strip() for p in passages) and settings.ocr_enabled:
            # Optional deployment extras: pip install pypdfium2 pytesseract; install tesseract OS package.
            import pypdfium2 as pdfium
            import pytesseract
            pdf = pdfium.PdfDocument(data)
            ocr_truncated = len(pdf) > 30
            passages = [{'reference': f'page {i + 1} (OCR, verify accuracy)', 'text': pytesseract.image_to_string(pdf[i].render(scale=2).to_pil(), timeout=20)} for i in range(min(len(pdf), 30))]
    elif ext == '.docx':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(f.file_size for f in z.infolist()) > 40_000_000 or len(z.infolist()) > 2000:
                raise ValueError('DOCX expanded content exceeds the processing limit.')
            if 'word/document.xml' not in z.namelist():
                raise ValueError('The file is not a Word DOCX document.')
        doc = Document(io.BytesIO(data))
        heading = 'Document'
        for para in doc.paragraphs:
            if para.style.name.startswith('Heading'):
                heading = para.text
            if para.text.strip():
                passages.append({'reference': heading, 'text': para.text})
        for i, table in enumerate(doc.tables):
            passages.append({'reference': f'table {i + 1}', 'text': '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)})
    else:
        raise ValueError('Supported formats: PDF, DOCX, TXT and Markdown.')
    chunks = []
    total = 0
    for p in passages:
        for offset in range(0, len(p['text']), 1800):
            text = p['text'][offset:offset + 1800]
            total += len(text)
            if total > 1_000_000:
                return chunks, 'Extraction stopped at 1 million characters; split the document for complete coverage.'
            if text.strip():
                chunks.append({'reference': p['reference'], 'text': text})
    warning = None if chunks else 'No readable text extracted. This may be a scanned PDF; enable optional OCR or upload a text-based copy.'
    if ocr_truncated:
        warning = 'OCR processed only the first 30 pages. Split the PDF to include the remaining pages; verify OCR accuracy.'
    if ext == '.pdf' and any(not p['text'].strip() for p in passages) and chunks:
        warning = 'Some pages contain no readable text and were excluded. Supply OCR text for those pages.'
    return chunks, warning


def build_context(prompt, preferences, attachments, history, decisions):
    from .providers import solution_context_fits
    words = set(re.findall(r'\w{3,}', prompt.lower()))
    ranked = []
    for a in attachments:
        for index, passage in enumerate(a.passages):
            score = sum(1 for word in words if word in passage['text'].lower())
            ranked.append((score, a.filename, index, passage, getattr(a, 'id', 'reference')))
    ranked.sort(key=lambda r: (-r[0], r[1], r[2]))
    base = {'current_prompt': prompt, 'preferences': preferences, 'approved_decisions': decisions,
            'history': [{'prompt': v.prompt[:3000], 'summary': (v.response or {}).get('summary', '')[:4000]} for v in history[-4:]],
            'reference_passages': [], 'document_warnings': [f'{a.filename}: {a.warning}' for a in attachments if a.warning]}
    # Account for the exact schema, system message, JSON escaping, separators,
    # provenance IDs and retry instructions, rather than assuming a fixed reserve.
    # Use a conservative retrieval note during selection, then update its count.
    total = len(ranked)
    base['retrieval_note'] = f'Included {total} of {total} passages using lexical relevance. Excluded passages have not been reviewed.'
    while base['history'] and not solution_context_fits(base):
        base['history'].pop(0)
    if not solution_context_fits(base):
        raise ValueError('Prompt and preferences exceed the configured input budget. Shorten the prompt or increase MAX_INPUT_TOKENS.')
    for _, filename, index, p, attachment_id in ranked:
        source = {'id': f'{filename}::{attachment_id}::{p["reference"]}::chunk-{index + 1}', 'filename': filename, **p}
        base['reference_passages'].append(source)
        if not solution_context_fits(base):
            base['reference_passages'].pop()
    base['retrieval_note'] = f'Included {len(base["reference_passages"])} of {len(ranked)} passages using lexical relevance. Excluded passages have not been reviewed.'
    return base
